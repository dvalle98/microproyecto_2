"""Extracción de documentos y transcripción de audio para la demo ODS."""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
import re
import wave


MAX_DOCUMENT_BYTES = 5 * 1024 * 1024
MAX_TEXT_CHARACTERS = 20_000
MAX_AUDIO_BYTES = 8 * 1024 * 1024
MAX_AUDIO_SECONDS = 60
SUPPORTED_DOCUMENT_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}


class UserInputError(ValueError):
    """Error de entrada que puede mostrarse directamente a la persona usuaria."""


@dataclass(frozen=True)
class ExtractedDocument:
    """Texto utilizable y metadatos mínimos de la extracción."""

    text: str
    original_characters: int
    truncated: bool


def _normalize_text(text: str) -> str:
    """Limpia artefactos básicos sin alterar el contenido lingüístico."""
    text = text.replace("\x00", "")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def _decode_plain_text(payload: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return payload.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise UserInputError("No fue posible reconocer la codificación del archivo de texto.")


def _extract_pdf(payload: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(BytesIO(payload))
        if reader.is_encrypted and not reader.decrypt(""):
            raise UserInputError("El PDF está protegido con contraseña y no puede leerse.")
        pages = [(page.extract_text() or "") for page in reader.pages]
    except UserInputError:
        raise
    except Exception as exc:
        raise UserInputError("El PDF no pudo leerse. Verifica que el archivo no esté dañado.") from exc
    return "\n\n".join(pages)


def _extract_docx(payload: bytes) -> str:
    from docx import Document

    try:
        document = Document(BytesIO(payload))
    except Exception as exc:
        raise UserInputError("El documento Word no pudo leerse. Verifica que sea un archivo .docx válido.") from exc

    blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return "\n\n".join(blocks)


def extract_document_text(filename: str, payload: bytes) -> ExtractedDocument:
    """Extrae texto de un archivo permitido y limita el contenido enviado al modelo."""
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_DOCUMENT_EXTENSIONS:
        raise UserInputError("Formato no admitido. Usa TXT, MD, PDF o DOCX.")
    if not payload:
        raise UserInputError("El archivo está vacío.")
    if len(payload) > MAX_DOCUMENT_BYTES:
        raise UserInputError("El archivo supera el límite de 5 MB.")

    if suffix in {".txt", ".md"}:
        raw_text = _decode_plain_text(payload)
    elif suffix == ".pdf":
        raw_text = _extract_pdf(payload)
    else:
        raw_text = _extract_docx(payload)

    text = _normalize_text(raw_text)
    if not text:
        hint = " Si es un PDF escaneado, primero necesita reconocimiento OCR." if suffix == ".pdf" else ""
        raise UserInputError(f"No se encontró texto utilizable en el archivo.{hint}")

    original_characters = len(text)
    truncated = original_characters > MAX_TEXT_CHARACTERS
    if truncated:
        text = text[:MAX_TEXT_CHARACTERS].rsplit(" ", 1)[0].rstrip()

    return ExtractedDocument(
        text=text,
        original_characters=original_characters,
        truncated=truncated,
    )


def audio_duration_seconds(payload: bytes) -> float:
    """Valida un WAV y devuelve su duración."""
    if not payload:
        raise UserInputError("La grabación está vacía.")
    if len(payload) > MAX_AUDIO_BYTES:
        raise UserInputError("La grabación supera el tamaño permitido. Intenta un dictado más breve.")
    try:
        with wave.open(BytesIO(payload), "rb") as audio:
            rate = audio.getframerate()
            frames = audio.getnframes()
    except (wave.Error, EOFError) as exc:
        raise UserInputError("La grabación no tiene un formato de audio válido.") from exc
    if not rate:
        raise UserInputError("No fue posible determinar la duración de la grabación.")
    return frames / rate


def transcribe_audio(payload: bytes, api_key: str) -> str:
    """Transcribe un WAV en español mediante la API de OpenAI."""
    duration = audio_duration_seconds(payload)
    if duration < 1:
        raise UserInputError("La grabación es demasiado breve. Dicta al menos una oración.")
    if duration > MAX_AUDIO_SECONDS:
        raise UserInputError("La grabación supera 60 segundos. Divide el dictado en una parte más breve.")
    if not api_key:
        raise UserInputError("La transcripción no está configurada en esta aplicación.")

    from openai import OpenAI

    client = OpenAI(api_key=api_key, timeout=45.0)
    response = client.audio.transcriptions.create(
        model="gpt-4o-mini-transcribe",
        file=("dictado.wav", payload, "audio/wav"),
        language="es",
        prompt="Transcribe fielmente este dictado en español sobre desarrollo sostenible y Objetivos de Desarrollo Sostenible.",
    )
    text = _normalize_text(response.text)
    if not text:
        raise UserInputError("No se reconoció voz en la grabación. Intenta nuevamente en un lugar más silencioso.")
    return text[:MAX_TEXT_CHARACTERS]
